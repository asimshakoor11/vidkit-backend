"""PDF conversion tools: Word, Excel, JPG."""

from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path

import img2pdf
from openpyxl import Workbook

from app.core.config import get_settings
from app.core.errors import AppError
from app.core.logging import get_logger
from app.services.pdf.bins import poppler_path, soffice_bin
from app.services.pdf.io import ProgressCb, out_name

logger = get_logger(__name__)


def _run_soffice(input_path: Path, out_dir: Path, convert_to: str) -> Path:
    """Run LibreOffice headless conversion; return the produced file."""
    soffice = soffice_bin()
    settings = get_settings()
    out_dir.mkdir(parents=True, exist_ok=True)

    # LibreOffice writes next to --outdir using the input stem
    with tempfile.TemporaryDirectory(dir=out_dir) as tmp:
        tmp_path = Path(tmp)
        # Copy input so LO never touches the upload path naming
        work_in = tmp_path / input_path.name
        shutil.copy2(input_path, work_in)
        cmd = [
            soffice,
            "--headless",
            "--nologo",
            "--nofirststartwizard",
            "--norestore",
            "--convert-to",
            convert_to,
            "--outdir",
            str(tmp_path),
            str(work_in),
        ]
        try:
            proc = subprocess.run(
                cmd,
                capture_output=True,
                text=True,
                timeout=settings.pdf_convert_timeout_sec,
                check=False,
            )
        except subprocess.TimeoutExpired as exc:
            raise AppError("CONVERSION_FAILED", "Document conversion timed out.") from exc
        except FileNotFoundError as exc:
            raise AppError("DEPENDENCY_MISSING", "LibreOffice is not installed.") from exc

        if proc.returncode != 0:
            logger.warning("soffice_failed", stderr=(proc.stderr or "")[-500:])
            raise AppError("CONVERSION_FAILED", "Document conversion failed.")

        # Find converted file
        candidates = list(tmp_path.glob(f"*.{convert_to.split(':')[0]}"))
        if not candidates:
            # Some LO builds use different extensions
            candidates = [p for p in tmp_path.iterdir() if p.is_file() and p != work_in]
        if not candidates:
            raise AppError("CONVERSION_FAILED", "Document conversion produced no output.")

        produced = candidates[0]
        final = out_dir / produced.name
        shutil.move(str(produced), str(final))
        return final


def pdf_to_word(
    input_path: Path,
    output_dir: Path,
    progress_callback: ProgressCb = None,
) -> Path:
    """Convert PDF to DOCX via LibreOffice."""
    if progress_callback:
        progress_callback(10, "converting")
    result = _run_soffice(input_path, output_dir, "docx")
    # Normalize name
    dest = output_dir / out_name(input_path.stem, "word", ".docx")
    if result != dest:
        if dest.exists():
            dest.unlink()
        result.rename(dest)
    if progress_callback:
        progress_callback(100, "done")
    return dest


def word_to_pdf(
    input_path: Path,
    output_dir: Path,
    progress_callback: ProgressCb = None,
) -> Path:
    """Convert DOCX to PDF via LibreOffice; light-validate with python-docx."""
    try:
        from docx import Document

        Document(str(input_path))
    except Exception as exc:  # noqa: BLE001
        raise AppError("UNSUPPORTED_FORMAT", "Could not read this Word document.") from exc

    if progress_callback:
        progress_callback(10, "converting")
    result = _run_soffice(input_path, output_dir, "pdf")
    dest = output_dir / out_name(input_path.stem, "from_word")
    if result != dest:
        if dest.exists():
            dest.unlink()
        result.rename(dest)
    if progress_callback:
        progress_callback(100, "done")
    return dest


def pdf_to_excel(
    input_path: Path,
    output_dir: Path,
    progress_callback: ProgressCb = None,
) -> Path:
    """Extract tables from PDF into an XLSX workbook."""
    import pdfplumber

    if progress_callback:
        progress_callback(10, "extracting")

    wb = Workbook()
    # Remove default sheet; recreate per table/page
    default = wb.active
    sheet_index = 0

    try:
        with pdfplumber.open(str(input_path)) as pdf:
            total = max(len(pdf.pages), 1)
            for page_i, page in enumerate(pdf.pages, start=1):
                tables = page.extract_tables() or []
                if not tables:
                    # Fallback: one column of text lines
                    text = page.extract_text() or ""
                    lines = [ln for ln in text.splitlines() if ln.strip()]
                    if not lines:
                        continue
                    sheet_index += 1
                    ws = wb.create_sheet(title=f"Page{page_i}"[:31])
                    for row_i, line in enumerate(lines, start=1):
                        ws.cell(row=row_i, column=1, value=line)
                else:
                    for t_i, table in enumerate(tables, start=1):
                        sheet_index += 1
                        name = f"P{page_i}T{t_i}"[:31]
                        ws = wb.create_sheet(title=name)
                        for r_i, row in enumerate(table or [], start=1):
                            for c_i, cell in enumerate(row or [], start=1):
                                ws.cell(row=r_i, column=c_i, value=cell)
                if progress_callback:
                    progress_callback(10 + page_i / total * 80, "extracting")
    except AppError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise AppError("CONVERSION_FAILED", "Failed to extract tables from PDF.") from exc

    if sheet_index == 0:
        ws = default
        ws.title = "Sheet1"
        ws.cell(row=1, column=1, value="No tables or text found in this PDF.")
    else:
        wb.remove(default)

    out = output_dir / out_name(input_path.stem, "excel", ".xlsx")
    output_dir.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    if progress_callback:
        progress_callback(100, "done")
    return out


def pdf_to_jpg(
    input_path: Path,
    output_dir: Path,
    dpi: int = 150,
    progress_callback: ProgressCb = None,
) -> list[Path]:
    """Render PDF pages to JPG via pdf2image (Poppler)."""
    from pdf2image import convert_from_path

    if progress_callback:
        progress_callback(10, "rendering")

    kwargs: dict = {"dpi": dpi, "fmt": "jpeg"}
    poppler = poppler_path()
    if poppler:
        kwargs["poppler_path"] = poppler

    try:
        images = convert_from_path(str(input_path), **kwargs)
    except Exception as exc:  # noqa: BLE001
        logger.warning("pdf2image_failed", error=str(exc))
        raise AppError(
            "CONVERSION_FAILED",
            "PDF to JPG failed. Ensure Poppler is installed and on PATH.",
        ) from exc

    if not images:
        raise AppError("UNSUPPORTED_FORMAT", "This PDF has no pages.")

    outputs: list[Path] = []
    stem = input_path.stem
    for i, img in enumerate(images, start=1):
        out = output_dir / out_name(stem, f"p{i:02d}", ".jpg")
        img.convert("RGB").save(str(out), "JPEG", quality=90)
        outputs.append(out)
        if progress_callback:
            progress_callback(10 + i / len(images) * 85, "rendering")

    if progress_callback:
        progress_callback(100, "done")
    return outputs


def jpg_to_pdf(
    paths: list[Path],
    output_dir: Path,
    progress_callback: ProgressCb = None,
) -> Path:
    """Combine images into a single PDF via img2pdf."""
    if not paths:
        raise AppError("UNSUPPORTED_FORMAT", "Upload at least one image.")

    if progress_callback:
        progress_callback(10, "converting")

    stem = paths[0].stem
    output_dir.mkdir(parents=True, exist_ok=True)
    out = output_dir / out_name(stem, "from_jpg")
    try:
        with out.open("wb") as fh:
            fh.write(img2pdf.convert([str(p) for p in paths]))
    except Exception as exc:  # noqa: BLE001
        raise AppError("CONVERSION_FAILED", "Could not convert images to PDF.") from exc

    if progress_callback:
        progress_callback(100, "done")
    return out
